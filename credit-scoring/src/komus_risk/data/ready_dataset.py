"""Универсальная загрузка готового табличного датасета без ML-обработки."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
from typing import Any

import pandas as pd

from komus_risk.contracts import DatasetContract
from komus_risk.hashing import stable_hash


_SUPPORTED_FORMATS = {".csv": "csv", ".xlsx": "xlsx", ".xlsb": "xlsb", ".parquet": "parquet"}


@dataclass(frozen=True, slots=True)
class LoadedDataset:
    """Загруженный dataframe и его метаданные.

    Контейнер неизменяем, однако сам dataframe намеренно не копируется и остаётся
    исходным результатом pandas-загрузки.
    """

    dataframe: pd.DataFrame
    contract: DatasetContract
    source_path: Path
    source_format: str
    source_file_sha256: str


class ReadyDatasetAdapter:
    """Загружает готовый файл и формирует валидированный паспорт датасета."""

    def load(
        self,
        path: str | Path,
        *,
        dataset_id: str,
        dataset_version: str,
        dataset_name: str,
        target_column: str,
        positive_class: str | int | float | bool,
        identifier_column: str,
        feature_registry_id: str,
        feature_registry_hash: str,
        final_test_locked: bool,
        separator: str = ",",
        encoding: str = "utf-8",
        sheet_name: str | int = 0,
    ) -> LoadedDataset:
        """Читает один готовый dataset без изменения его строк или колонок."""
        source_path = Path(path)
        self._validate_source_path(source_path)
        source_format = self._source_format(source_path)
        self._validate_source_headers(
            source_path,
            source_format,
            separator=separator,
            encoding=encoding,
            sheet_name=sheet_name,
        )
        dataframe = self._read(
            source_path,
            source_format,
            separator=separator,
            encoding=encoding,
            sheet_name=sheet_name,
        )
        self._validate_dataframe(
            dataframe,
            target_column=target_column,
            positive_class=positive_class,
            identifier_column=identifier_column,
        )

        source_file_sha256 = self._file_sha256(source_path)
        dataset_fingerprint = stable_hash(
            {
                "source_file_sha256": source_file_sha256,
                "source_format": source_format,
                "read_options": self._fingerprint_options(
                    source_format,
                    separator=separator,
                    encoding=encoding,
                    sheet_name=sheet_name,
                ),
            }
        )
        contract = DatasetContract(
            dataset_id=dataset_id,
            dataset_version=dataset_version,
            dataset_name=dataset_name,
            source_type=f"ready_{source_format}",
            dataset_fingerprint=dataset_fingerprint,
            row_count=len(dataframe.index),
            column_count=len(dataframe.columns),
            target_column=target_column,
            positive_class=positive_class,
            identifier_column=identifier_column,
            feature_registry_id=feature_registry_id,
            feature_registry_hash=feature_registry_hash,
            validation_status="validated",
            final_test_locked=final_test_locked,
        )
        return LoadedDataset(
            dataframe=dataframe,
            contract=contract,
            source_path=source_path,
            source_format=source_format,
            source_file_sha256=source_file_sha256,
        )

    @staticmethod
    def _validate_source_path(source_path: Path) -> None:
        if not source_path.exists():
            raise FileNotFoundError(f"Файл датасета не найден: «{source_path}».")
        if not source_path.is_file():
            raise ValueError(f"Путь к датасету должен указывать на файл: «{source_path}».")

    @staticmethod
    def _source_format(source_path: Path) -> str:
        try:
            return _SUPPORTED_FORMATS[source_path.suffix.lower()]
        except KeyError as error:
            supported = ", ".join(sorted(_SUPPORTED_FORMATS))
            raise ValueError(
                f"Неподдерживаемое расширение «{source_path.suffix}». Поддерживаются: {supported}."
            ) from error

    @classmethod
    def _validate_source_headers(
        cls,
        source_path: Path,
        source_format: str,
        *,
        separator: str,
        encoding: str,
        sheet_name: str | int,
    ) -> None:
        if source_format == "csv":
            headers = cls._read_csv_header(source_path, separator=separator, encoding=encoding)
        elif source_format == "xlsx":
            headers = cls._read_xlsx_header(source_path, sheet_name=sheet_name)
        elif source_format == "xlsb":
            headers = cls._read_xlsb_header(source_path, sheet_name=sheet_name)
        else:
            return
        cls._raise_duplicate_header_error(headers)

    @staticmethod
    def _read_csv_header(source_path: Path, *, separator: str, encoding: str) -> list[str]:
        if len(separator) != 1:
            raise ValueError("Разделитель CSV должен состоять из одного символа.")
        with source_path.open("r", encoding=encoding, newline="") as source_file:
            return next(csv.reader(source_file, delimiter=separator), [])

    @staticmethod
    def _read_xlsx_header(source_path: Path, *, sheet_name: str | int) -> list[object]:
        from openpyxl import load_workbook

        workbook = load_workbook(source_path, read_only=True, data_only=False)
        try:
            worksheet = workbook.worksheets[sheet_name] if isinstance(sheet_name, int) else workbook[sheet_name]
            return list(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
        finally:
            workbook.close()

    @staticmethod
    def _read_xlsb_header(source_path: Path, *, sheet_name: str | int) -> list[object]:
        from pyxlsb import open_workbook

        with open_workbook(source_path) as workbook:
            selected_sheet = sheet_name + 1 if isinstance(sheet_name, int) else sheet_name
            with workbook.get_sheet(selected_sheet) as worksheet:
                return [cell.v for cell in next(worksheet.rows(), ())]

    @staticmethod
    def _raise_duplicate_header_error(headers: list[object]) -> None:
        duplicate_headers = []
        seen_headers = []
        for header in headers:
            if header in seen_headers and header not in duplicate_headers:
                duplicate_headers.append(header)
            seen_headers.append(header)
        if duplicate_headers:
            names = ", ".join(f"«{header}»" for header in duplicate_headers)
            raise ValueError(f"Исходный файл содержит повторяющиеся имена колонок: {names}.")

    @staticmethod
    def _read(
        source_path: Path,
        source_format: str,
        *,
        separator: str,
        encoding: str,
        sheet_name: str | int,
    ) -> pd.DataFrame:
        if source_format == "csv":
            return pd.read_csv(source_path, sep=separator, encoding=encoding)
        if source_format == "xlsx":
            return pd.read_excel(source_path, sheet_name=sheet_name, engine="openpyxl")
        if source_format == "xlsb":
            return pd.read_excel(source_path, sheet_name=sheet_name, engine="pyxlsb")
        return pd.read_parquet(source_path)

    @staticmethod
    def _fingerprint_options(
        source_format: str,
        *,
        separator: str,
        encoding: str,
        sheet_name: str | int,
    ) -> dict[str, Any]:
        if source_format == "csv":
            return {"separator": separator, "encoding": encoding}
        if source_format in {"xlsx", "xlsb"}:
            return {"sheet_name": sheet_name}
        return {}

    @staticmethod
    def _validate_dataframe(
        dataframe: pd.DataFrame,
        *,
        target_column: str,
        positive_class: str | int | float | bool,
        identifier_column: str,
    ) -> None:
        if dataframe.empty:
            raise ValueError("Датасет пуст: требуется хотя бы одна строка и один столбец.")
        if not dataframe.columns.is_unique:
            raise ValueError("Имена колонок датасета должны быть однозначными.")
        if target_column not in dataframe.columns:
            raise ValueError(f"Целевая колонка «{target_column}» отсутствует в датасете.")
        if identifier_column not in dataframe.columns:
            raise ValueError(f"Колонка идентификатора «{identifier_column}» отсутствует в датасете.")
        if target_column == identifier_column:
            raise ValueError("Целевая колонка и колонка идентификатора должны различаться.")

        target = dataframe[target_column]
        if target.isna().any():
            raise ValueError("Целевая колонка не должна содержать пропущенные значения.")
        if not target.eq(positive_class).any():
            raise ValueError(
                f"Положительный класс «{positive_class}» отсутствует в целевой колонке."
            )

    @staticmethod
    def _file_sha256(source_path: Path) -> str:
        digest = sha256()
        with source_path.open("rb") as source_file:
            for chunk in iter(lambda: source_file.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()
