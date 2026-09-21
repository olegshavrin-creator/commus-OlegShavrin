"""Physical, semantics-free reading of supported tabular files."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
from typing import Any

import pandas as pd

from komus_risk.hashing import stable_hash

_FORMATS = {".csv": "csv", ".xlsx": "xlsx", ".xlsb": "xlsb", ".parquet": "parquet"}


class TabularReadError(ValueError):
    """A stable, user-displayable physical source error."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True, slots=True)
class TabularSnapshot:
    source_path: Path
    source_format: str
    read_options: dict[str, Any]
    source_file_sha256: str
    fingerprint: str
    row_count: int
    column_count: int
    dataframe: pd.DataFrame
    physical_headers: tuple[str, ...] = ()


class TabularReader:
    """Reads bytes into a dataframe without assigning dataset semantics."""

    def read(self, path: str | Path, *, separator: str = ",", encoding: str = "utf-8", sheet_name: str | int = 0) -> TabularSnapshot:
        source_path = Path(path)
        if not source_path.exists():
            raise TabularReadError("file_not_found", f"Файл не найден: «{source_path}».")
        if not source_path.is_file():
            raise TabularReadError("path_not_file", f"Путь не указывает на файл: «{source_path}».")
        try:
            source_format = _FORMATS[source_path.suffix.lower()]
        except KeyError as error:
            raise TabularReadError("unsupported_format", f"Неподдерживаемый формат: «{source_path.suffix}».") from error
        if source_format == "csv" and len(separator) != 1:
            raise TabularReadError("invalid_read_options", "Разделитель CSV должен состоять из одного символа.")
        options = self._options(source_format, separator, encoding, sheet_name)
        physical_headers = self._physical_headers(source_path, source_format, separator, encoding, sheet_name)
        try:
            dataframe = self._read(source_path, source_format, separator, encoding, sheet_name)
        except TabularReadError:
            raise
        except Exception as error:
            raise TabularReadError("unreadable_source", f"Не удалось прочитать табличный файл: {error}") from error
        if tuple(dataframe.columns) != physical_headers:
            raise TabularReadError("header_identity_mismatch", "Pandas columns do not match physical source headers.")
        file_hash = self._sha256(source_path)
        return TabularSnapshot(source_path, source_format, options, file_hash,
            stable_hash({"source_file_sha256": file_hash, "source_format": source_format, "read_options": options}),
            len(dataframe.index), len(dataframe.columns), dataframe, physical_headers)

    @staticmethod
    def _options(fmt: str, separator: str, encoding: str, sheet_name: str | int) -> dict[str, Any]:
        if fmt == "csv": return {"separator": separator, "encoding": encoding}
        if fmt in {"xlsx", "xlsb"}: return {"sheet_name": sheet_name}
        return {}

    @staticmethod
    def _read(path: Path, fmt: str, separator: str, encoding: str, sheet_name: str | int) -> pd.DataFrame:
        if fmt == "csv": return pd.read_csv(path, sep=separator, encoding=encoding)
        if fmt == "xlsx": return pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl")
        if fmt == "xlsb": return pd.read_excel(path, sheet_name=sheet_name, engine="pyxlsb")
        return pd.read_parquet(path)

    @classmethod
    def _physical_headers(cls, path: Path, fmt: str, separator: str, encoding: str, sheet_name: str | int) -> tuple[str, ...]:
        try:
            if fmt == "csv":
                with path.open("r", encoding=encoding, newline="") as file:
                    headers = next(csv.reader(file, delimiter=separator), [])
            elif fmt == "xlsx":
                from openpyxl import load_workbook
                book = load_workbook(path, read_only=True, data_only=False)
                try:
                    sheet = book.worksheets[sheet_name] if isinstance(sheet_name, int) else book[sheet_name]
                    headers = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
                finally: book.close()
            elif fmt == "xlsb":
                from pyxlsb import open_workbook
                with open_workbook(path) as book:
                    with book.get_sheet(sheet_name + 1 if isinstance(sheet_name, int) else sheet_name) as sheet:
                        headers = [cell.v for cell in next(sheet.rows(), ())]
            else:
                import pyarrow.parquet as pq
                headers = list(pq.ParquetFile(path).schema_arrow.names)
        except Exception as error:
            raise TabularReadError("invalid_read_options", f"Не удалось прочитать заголовки: {error}") from error
        if any(not isinstance(header, str) or not header.strip() for header in headers):
            raise TabularReadError("invalid_physical_header", "Physical headers must be non-empty strings.")
        if len(headers) != len(set(headers)):
            raise TabularReadError("invalid_physical_header", "Physical headers must be unique.")
        return tuple(headers)

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = sha256()
        with path.open("rb") as file:
            for chunk in iter(lambda: file.read(1024 * 1024), b""): digest.update(chunk)
        return digest.hexdigest()
